import csv
import io
import re
from pathlib import Path


ALIASES = {
    'ref':['ref','designator','reference','位置','位号'],
    'name':['name','子项物料名称','物料名称'],
    'spec':['spec','子项规格型号','规格型号'],
    'partno':['partno','物料编号','料号'],
    'qty':['qty','quantity','数量'],
    'footprint':['footprint','package','封装'],
    'x':['x','center-x(mm)','mid x','x坐标'],
    'y':['y','center-y(mm)','mid y','y坐标'],
    'rotation':['rotation','angle','角度'],
    'side':['layer','side','面别'],
}


def read_table(filename, data):
    suffix = Path(filename).suffix.lower()
    encoding = 'workbook'
    if suffix in ('.xlsx','.xls'):
        if suffix == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                rows = [[str(v) if v is not None else '' for v in row]
                        for row in book.active.iter_rows(values_only=True)]
            finally:
                book.close()
        else:
            import xlrd
            sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
            rows = [[str(v) for v in sheet.row_values(i)] for i in range(sheet.nrows)]
    elif suffix in ('.csv','.txt','.tsv','.cad','.bom'):
        for encoding in ('utf-8-sig','gb18030','utf-16'):
            try:
                content = data.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            raise ValueError('encoding: 无法解码，请保存为 UTF-8 或 GBK')
        lines = content.splitlines()
        sample = '\n'.join(lines)
        if '\t' in sample:
            delimiter='\t'
        else:
            delimiter=',' if sample.count(',')>=sample.count(';') else ';'
        rows=list(csv.reader(io.StringIO(content),delimiter=delimiter))
        if max((len(r) for r in rows),default=0)==1:
            rows=[re.split(r'\s{2,}|\t',r[0].strip()) if r else [] for r in rows]
    else:
        raise ValueError('format: 支持 CSV/TXT/TSV/XLSX/XLS')
    if not rows:
        raise ValueError('empty_table: 文件没有数据')
    if len(rows)>100000:
        raise ValueError('table_size: 请将表格限制在 100000 行内')
    scores=[sum(any(str(c).strip().lower() in names for names in ALIASES.values()) for c in row)
            for row in rows[:100]]
    header=max(range(len(scores)),key=scores.__getitem__) if max(scores)>0 else -1
    mapping={}
    if header>=0:
        for field,names in ALIASES.items():
            matches=[i for i,c in enumerate(rows[header]) if c.strip().lower() in names]
            if len(matches)==1:
                mapping[field]=matches[0]
    return {'filename':Path(filename).name,'encoding':encoding,'header_row':header,
            'mapping':mapping,'excluded_rows':[],
            'rows':[{'id':i,'cells':row} for i,row in enumerate(rows)]}


def references(value):
    result=[]
    for token in re.split(r'[,;\s，；、]+',str(value).strip().upper()):
        if not token:
            continue
        match=re.fullmatch(r'([A-Z]+)(\d+)\s*[-~]\s*([A-Z]*)(\d+)',token)
        if match:
            prefix,a,other,b=match.groups()
            if other and other!=prefix:
                raise ValueError('ref_range: 位号范围前缀不一致')
            if int(b)<int(a) or int(b)-int(a)>10000:
                raise ValueError('ref_range: 位号范围无效')
            result.extend(f'{prefix}{n}' for n in range(int(a),int(b)+1))
        elif re.fullmatch(r'[A-Z][A-Z0-9_.-]*\d[A-Z0-9_.-]*',token):
            result.append(token)
        else:
            raise ValueError(f'invalid_ref: {token}')
    return result


def number(value):
    text=str(value).strip()
    match=re.fullmatch(r'([-+]?\d*\.?\d+(?:[eE][-+]?\d+)?)\s*(?:mm|mil|inch|in|°|deg)?',text,re.I)
    if not match:
        raise ValueError(f'invalid_number: {text}')
    import math
    n=float(match.group(1))
    if not math.isfinite(n):
        raise ValueError('invalid_number: 必须是有限数值')
    return n


def side(value):
    text=str(value).strip().lower()
    if text in ('bottom','bottomlayer','b','bot','背面','底面'):return 'bottom'
    if text in ('top','toplayer','t','正面','顶面'):return 'top'
    raise ValueError(f'invalid_side: {value}')


def cell(row,mapping,field,default=''):
    index=mapping.get(field)
    return row[index].strip() if index is not None and index<len(row) else default
